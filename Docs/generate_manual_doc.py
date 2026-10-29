# -*- coding: utf-8 -*-
import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_table_styling(table, col_widths=None, header_bg="1F3A60", alt_bg="F1F5F9"):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            if col_widths and idx < len(col_widths):
                cell.width = Inches(col_widths[idx])
                
    # Header row
    for cell in table.rows[0].cells:
        set_cell_background(cell, header_bg)
        for p in cell.paragraphs:
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            for run in p.runs:
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.font.size = Pt(9.5)
                run.font.name = "Segoe UI"
                
    # Data rows
    for r_idx, row in enumerate(table.rows[1:], start=1):
        bg = alt_bg if r_idx % 2 == 1 else "FFFFFF"
        for cell in row.cells:
            set_cell_background(cell, bg)
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                for run in p.runs:
                    run.font.size = Pt(9)
                    run.font.name = "Segoe UI"

def add_heading_with_spacing(doc, text, level):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    for run in p.runs:
        run.font.name = "Segoe UI"
    return p

def add_p(doc, text, bold_prefix=None, style='Normal'):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r1 = p.add_run(bold_prefix)
        r1.font.bold = True
        r1.font.name = "Segoe UI"
        r1.font.size = Pt(10)
    r2 = p.add_run(text)
    r2.font.name = "Segoe UI"
    r2.font.size = Pt(10)
    return p

def add_bullet(doc, title, text):
    p = doc.add_paragraph(style='List Paragraph')
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    r1 = p.add_run(title + ": ")
    r1.font.bold = True
    r1.font.name = "Segoe UI"
    r1.font.size = Pt(10)
    r2 = p.add_run(text)
    r2.font.name = "Segoe UI"
    r2.font.size = Pt(10)
    return p

def add_table_data(doc, headers, data, col_widths=None):
    table = doc.add_table(rows=len(data) + 1, cols=len(headers))
    table.style = 'Table Grid'
    for c_idx, h in enumerate(headers):
        table.cell(0, c_idx).text = h
    for r_idx, row in enumerate(data, start=1):
        for c_idx, val in enumerate(row):
            table.cell(r_idx, c_idx).text = str(val)
    set_table_styling(table, col_widths=col_widths)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(6)
    return table

def generate_manual_doc(filepath):
    doc = docx.Document()
    
    # Title
    p_title = doc.add_heading("VerbFlow — Manual de Usuario y Guía de Operación Integral", level=1)
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(2)
    
    p_meta = doc.add_paragraph()
    p_meta.paragraph_format.space_after = Pt(12)
    r_meta = p_meta.add_run("Versión 1.0 (Guía de Roles) · Estudiantes, Docentes y Owner · I.E. Promoción Social (Palermo, Huila) · Octubre de 2026")
    r_meta.font.italic = True
    r_meta.font.size = Pt(9.5)
    r_meta.font.color.rgb = RGBColor(100, 116, 139)
    
    # 1. Introducción
    add_heading_with_spacing(doc, "1. Introducción a la Plataforma VerbFlow", level=2)
    add_p(doc, "VerbFlow es la plataforma web interactiva y privada diseñada especialmente para la Institución Educativa Promoción Social (Palermo, Huila), orientada a fortalecer las competencias lingüísticas en los tiempos verbales del idioma inglés para los estudiantes de educación media (grados 10° y 11°). Este manual detalla los procedimientos de uso, navegación y administración para cada uno de los roles autorizados en el sistema.")
    
    # 2. Guía del Estudiante
    add_heading_with_spacing(doc, "2. Módulo del Estudiante (Student Guide)", level=2)
    add_p(doc, "El portal del estudiante ofrece un entorno dinámico, táctil y libre de presiones innecesarias de gamificación, estructurado en cuatro secciones principales:")
    
    add_heading_with_spacing(doc, "2.1 Registro e Inicio de Sesión", level=3)
    add_bullet(doc, "Creación de Cuenta", "Ingresa a la pantalla principal, selecciona 'Registrarse como Estudiante', digita tu nombre completo, un nickname institucional único (ej. 'kevin10'), una contraseña segura y tu grado (10° u 11°).")
    add_bullet(doc, "Consentimiento Legal (Ley 1581)", "Marca obligatoriamente la casilla de autorización de datos personales y confirmación de consentimiento de acudiente.")
    add_bullet(doc, "Inicio de Sesión", "Ingresa con tu nickname y contraseña. El sistema mantendrá tu sesión abierta de forma segura.")

    add_heading_with_spacing(doc, "2.2 Modo Estudio y Chuleta de Fórmulas", level=3)
    add_bullet(doc, "Consulta Curricular", "Selecciona cualquiera de los 7 tiempos verbales en la barra lateral (Present Simple, Past Continuous, etc.) para ver su contexto de uso en español, estructuras sintácticas, reglas ortográficas y ejemplos ilustrativos.")
    add_bullet(doc, "Línea de Tiempo Visual", "Explora la posición cronológica de las acciones verbales respecto al momento del habla.")
    add_bullet(doc, "Entrenamiento Auditivo", "Presiona el ícono del altavoz junto a cada oración para escuchar la pronunciación nativa del navegador.")
    add_bullet(doc, "Chuleta Rápida de Fórmulas", "Haz clic en el botón flotante 'Chuleta' para desplegar en cualquier momento una ventana emergente con todas las fórmulas afirmativas, negativas e interrogativas resumidas.")

    add_heading_with_spacing(doc, "2.3 Modo Diálogos Contextuales y Práctica Oral con Micrófono", level=3)
    add_p(doc, "El modo conversacional te permite practicar el inglés en situaciones de la vida real con retroalimentación fonética inmediata:")
    add_bullet(doc, "Paso 1: Selección del Escenario", "Elige una situación comunicativa (ej. 'El Alibi del Detective' para practicar Past Continuous). Lee el contexto y el tono emocional de la escena.")
    add_bullet(doc, "Paso 2: Escucha del Interlocutor", "El personaje virtual emite su turno hablado. Puedes escucharlo a velocidad normal (1.0x) o en velocidad reducida (0.75x) para afinar el oído. Si lo requieres, activa los subtítulos en español.")
    add_bullet(doc, "Paso 3: Tu Turno de Habla (Speaking Challenge)", "La pantalla te indicará tu objetivo comunicativo en español (ej. 'Dile al detective que estabas viendo tele con tu familia'). Tienes la fórmula de apoyo visible.")
    add_bullet(doc, "Paso 4: Grabación con el Micrófono", "Mantén presionado el botón central 3D del micrófono o la barra espaciadora del teclado. Habla de forma clara. Observa cómo las ondas de audio reaccionan al volumen de tu voz.")
    add_bullet(doc, "Paso 5: Retroalimentación Fonética (Word-Diff)", "Al soltar el micrófono, el sistema pinta en verde las palabras correctas y en rojo las palabras que confundiste o pronunciaste mal.")
    add_bullet(doc, "Paso 6: Consejos del Coach Fonético", "Si cometiste un error (por ejemplo, decir 'washing' en lugar de 'watching'), lee el consejo en español sobre cómo posicionar la lengua y los labios, escucha la diferencia y vuelve a intentarlo.")

    add_heading_with_spacing(doc, "2.4 Práctica Libre de Gramática", level=3)
    add_bullet(doc, "Personalización", "Elige qué temas verbales practicar, el nivel de dificultad (Fácil, Media, Difícil o Ponderada) y el número de preguntas (5 a 25).")
    add_bullet(doc, "Ejercicios Mixtos", "Resuelve preguntas de selección múltiple con única respuesta y preguntas de completar escribiendo la forma verbal correcta.")
    add_bullet(doc, "Revisión Inmediata", "Al terminar, la web califica en tu propio navegador y te muestra la justificación de cada acierto o error. No se guarda nada en la base de datos.")

    add_heading_with_spacing(doc, "2.5 Presentación de Evaluaciones Oficiales de Clase (Modo Examen)", level=3)
    add_bullet(doc, "Ingreso con Código VF-XXXX", "Digita en la casilla de examen el código alfanumérico que te entregó tu profesor (ej. VF-4821).")
    add_bullet(doc, "Activación de Modo Examen", "Al pulsar 'Iniciar Examen', la pantalla entra en modo seguro: la teoría y la Chuleta de Fórmulas quedan bloqueadas temporalmente y se activa el temporizador regresivo.")
    add_bullet(doc, "Autoguardado Automático", "Cada respuesta que marcas o escribes se guarda de inmediato en el servidor. Si se desconecta el wifi o se apaga el equipo, no perderás lo que ya respondiste.")
    add_bullet(doc, "Entrega y Calificación Oficial (1.0 a 5.0)", "Al finalizar o cumplirse el tiempo límite, el examen se cierra y se computa tu calificación en la escala estándar de Colombia (1.0 a 5.0). Podrás ver tu nota y consultar tu historial académico.")

    # 3. Guía del Docente
    add_heading_with_spacing(doc, "3. Módulo del Docente (Teacher Suite)", level=2)
    add_p(doc, "El portal del docente (`teacher.html`) permite gestionar grupos escolares, crear evaluaciones protegidas y supervisar el desempeño del curso en tiempo real.")
    
    add_heading_with_spacing(doc, "3.1 Registro y Aprobación de Cuenta Docente", level=3)
    add_bullet(doc, "Registro Inicial", "Regístrate seleccionando el rol 'Docente'. La cuenta nacerá en estado 'pendiente' por seguridad institucional.")
    add_bullet(doc, "Activación por el Owner", "El administrador general (TheKeas) verificará tu vinculación a la institución y activará tu cuenta. Una vez aprobada, tendrás acceso pleno a la Teacher Suite.")

    add_heading_with_spacing(doc, "3.2 Administración de Clases y Estudiantes", level=3)
    add_bullet(doc, "Crear Nueva Clase", "Ingresa a 'Mis Clases' > 'Crear Clase'. Asigna un nombre al curso (ej. 'Inglés 10-A'), define el grado (10 u 11) y el año lectivo.")
    add_bullet(doc, "Regla de Oro (1 Estudiante = 1 Clase)", "Un estudiante solo puede pertenecer a un curso a la vez. Los alumnos quedan vinculados automáticamente a tu grupo al presentar su primera evaluación oficial con tu código.")
    add_bullet(doc, "Desvinculación", "Puedes retirar alumnos de tu lista si cambiaron de salón o fueron matriculados por error.")

    add_heading_with_spacing(doc, "3.3 Creación y Publicación de Evaluaciones", level=3)
    add_bullet(doc, "Configuración Pedagógica", "Selecciona los temas verbales a evaluar, la dificultad y la cantidad de preguntas (entre 5 y 50).")
    add_bullet(doc, "Parámetros de Aplicación", "Define fecha y hora exacta de apertura y cierre, duración máxima en minutos, intentos permitidos (por defecto 1) y modo de visualización de resultados (inmediato o diferido).")
    add_bullet(doc, "Generación del Código VF-XXXX", "Al guardar la prueba, el sistema genera automáticamente el código de acceso alfanumérico único para compartir con los estudiantes.")
    add_bullet(doc, "Publicación y Congelación de Preguntas", "Al hacer clic en 'Publicar', el sistema congela las preguntas en la base de datos para que todos los alumnos del grupo respondan el mismo examen sin discrepancias.")

    add_heading_with_spacing(doc, "3.4 Monitoreo en Vivo y Reporte de Calificaciones", level=3)
    add_bullet(doc, "Tablero en Tiempo Real", "Durante la aplicación de la prueba, observa en pantalla quién está rindiendo el examen, quién ya entregó, el tiempo transcurrido y la nota obtenida.")
    add_bullet(doc, "Alertas de Integridad de Foco", "El sistema te notifica si un estudiante cambió de pestaña o minimizó la pantalla durante la prueba.")
    add_bullet(doc, "Exportación a Excel / CSV", "Con un solo clic, descarga la planilla completa de calificaciones en formato Excel con nombres, aciertos y notas finales en escala 1.0 a 5.0, lista para cargar en el sistema de boletines del colegio.")

    add_heading_with_spacing(doc, "3.5 Aporte al Banco de Preguntas", level=3)
    add_bullet(doc, "Preguntas Propias", "Puedes redactar preguntas de selección múltiple o de completar. Estas preguntas se guardan con visibilidad 'privada' para tu uso exclusivo, salvo que el Owner las apruebe para todo el colegio.")

    # 4. Guía del Administrador (Owner TheKeas)
    add_heading_with_spacing(doc, "4. Módulo del Administrador (Owner TheKeas)", level=2)
    add_p(doc, "El Owner custodia el funcionamiento global y la seguridad de la plataforma a través de `owner.html`:")
    add_bullet(doc, "Aprobación de Profesores", "Revisa la lista de profesores en estado 'pendiente' y pulsa 'Aprobar' o 'Rechazar' con un solo clic. Esto impide que alumnos se hagan pasar por profesores.")
    add_bullet(doc, "Gestión de Usuarios", "Capacidad de suspender o reactivar cuentas ante faltas disciplinarias o directrices institucionales.")
    add_bullet(doc, "Supervisión del Banco Global", "Gestionar los 7 temas verbales, fórmulas, audios y publicar preguntas docentes al banco curricular general.")
    add_bullet(doc, "Auditoría del Sistema", "Consultar la bitácora de eventos sensibles con fechas, horas y direcciones IP.")

    # 5. Preguntas Frecuentes y Soporte Técnico (FAQ)
    add_heading_with_spacing(doc, "5. Preguntas Frecuentes y Soporte Técnico (FAQ)", level=2)
    h_faq = ["Problema o Duda Frecuente", "Solución Paso a Paso"]
    d_faq = [
        ["El navegador no detecta mi voz o bloquea el micrófono", "Haz clic en el ícono del candado junto a la URL del navegador en la barra de direcciones. Asegúrate de que el permiso de 'Micrófono' esté configurado en 'Permitir'. Si usas audífonos, verifica que el micrófono esté bien conectado."],
        ["Se cayó la conexión de internet mientras presentaba un examen", "Mantén la calma. VerbFlow guarda cada respuesta automáticamente en el servidor. Al restablecerse el wifi o los datos, el sistema sincroniza la última respuesta y te permite continuar hasta que finalice tu tiempo límite."],
        ["El código VF-XXXX arroja 'Examen fuera de horario'", "Verifica la fecha y hora de tu dispositivo. La prueba solo se activa dentro del rango de apertura programado por el docente."],
        ["Olvidé mi contraseña de acceso", "Si registraste tu correo electrónico opcional, usa el enlace 'Recuperar contraseña' en la pantalla de inicio; de lo contrario, solicita el restablecimiento directamente a tu docente o al Owner."],
        ["¿Por qué el sistema me marca error al decir 'watching'?", "Revisa el consejo del Coach Fonético. Asegúrate de marcar el sonido fuerte de la 'ch' /tʃ/ y no arrastrar el sonido suave /ʃ/ de 'washing'."]
    ]
    add_table_data(doc, h_faq, d_faq, col_widths=[2.8, 4.2])

    doc.save(filepath)
    print(f"Manual de Usuario guardado exitosamente en {filepath}")

generate_manual_doc("Docs/VerbFlow — Manual de Usuario y Operación.docx")
