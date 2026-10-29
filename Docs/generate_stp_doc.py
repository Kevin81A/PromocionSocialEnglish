# -*- coding: utf-8 -*-
import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_table_styling(table, col_widths=None, header_bg="1F3A60", alt_bg="F1F5F9"):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            if col_widths and idx < len(col_widths):
                cell.width = Inches(col_widths[idx])
                
    # Header row
    for cell in table.rows[0].cells:
        set_cell_background(cell, header_bg)
        for p in cell.paragraphs:
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            for run in p.runs:
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.font.size = Pt(9.5)
                run.font.name = "Segoe UI"
                
    # Data rows
    for r_idx, row in enumerate(table.rows[1:], start=1):
        bg = alt_bg if r_idx % 2 == 1 else "FFFFFF"
        for cell in row.cells:
            set_cell_background(cell, bg)
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                for run in p.runs:
                    run.font.size = Pt(9)
                    run.font.name = "Segoe UI"

def add_heading_with_spacing(doc, text, level):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    for run in p.runs:
        run.font.name = "Segoe UI"
    return p

def add_p(doc, text, bold_prefix=None, style='Normal'):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r1 = p.add_run(bold_prefix)
        r1.font.bold = True
        r1.font.name = "Segoe UI"
        r1.font.size = Pt(10)
    r2 = p.add_run(text)
    r2.font.name = "Segoe UI"
    r2.font.size = Pt(10)
    return p

def add_bullet(doc, title, text):
    p = doc.add_paragraph(style='List Paragraph')
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    r1 = p.add_run(title + ": ")
    r1.font.bold = True
    r1.font.name = "Segoe UI"
    r1.font.size = Pt(10)
    r2 = p.add_run(text)
    r2.font.name = "Segoe UI"
    r2.font.size = Pt(10)
    return p

def add_table_data(doc, headers, data, col_widths=None):
    table = doc.add_table(rows=len(data) + 1, cols=len(headers))
    table.style = 'Table Grid'
    for c_idx, h in enumerate(headers):
        table.cell(0, c_idx).text = h
    for r_idx, row in enumerate(data, start=1):
        for c_idx, val in enumerate(row):
            table.cell(r_idx, c_idx).text = str(val)
    set_table_styling(table, col_widths=col_widths)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(6)
    return table

def generate_stp_doc(filepath):
    doc = docx.Document()
    
    # Title
    p_title = doc.add_heading("VerbFlow — Plan de Pruebas de Software (STP) y Matriz de Casos de Prueba", level=1)
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(2)
    
    p_meta = doc.add_paragraph()
    p_meta.paragraph_format.space_after = Pt(12)
    r_meta = p_meta.add_run("Versión 1.0 (Definitiva) · Autor: Kevin · Basado en el Estándar IEEE 829:2008 · Octubre de 2026")
    r_meta.font.italic = True
    r_meta.font.size = Pt(9.5)
    r_meta.font.color.rgb = RGBColor(100, 116, 139)
    
    # 1. Introducción y Alcance
    add_heading_with_spacing(doc, "1. Introducción y Alcance", level=2)
    add_p(doc, "Este documento establece la estrategia formal de aseguramiento de calidad (QA), el plan de pruebas y la matriz de casos de prueba detallada para VerbFlow, plataforma educativa privada de la Institución Educativa Promoción Social (Palermo, Huila, Colombia). Su propósito es verificar y validar el cumplimiento riguroso de los requisitos funcionales del SRS v0.3, las 19 entidades del Modelo de Datos v0.3 y los 34 endpoints del Documento de Arquitectura y Especificación de la API (SAD) v1.0.")
    
    add_heading_with_spacing(doc, "1.1 Objetivos de Prueba", level=3)
    add_bullet(doc, "Integridad de Calificaciones", "Verificar que la fórmula matemática de 1.00 a 5.00 se compute con exactitud atómica en el servidor y que ninguna respuesta correcta sea accesible desde el navegador antes o durante una evaluación oficial.")
    add_bullet(doc, "Efectividad del Módulo Oral Comunicativo", "Validar el funcionamiento del micrófono mediante Web Speech API, la fidelidad del visualizador de ondas sonoras y la precisión del algoritmo Word-Diff al resaltar aciertos y desvíos fonéticos con feedback en español.")
    add_bullet(doc, "Resiliencia ante Conectividad Escolar Deficiente", "Comprobar que ante desconexiones repentinas de wifi o cortes eléctricos, las respuestas de los estudiantes se preserven íntegramente gracias al autoguardado progresivo.")
    add_bullet(doc, "Aislamiento por Rol y Docente", "Confirmar que los docentes requieran aprobación obligatoria del Owner y que un profesor solo pueda acceder y exportar las notas de sus propios cursos.")

    # 2. Estrategia y Niveles de Prueba
    add_heading_with_spacing(doc, "2. Estrategia y Niveles de Prueba", level=2)
    h_niveles = ["Nivel de Prueba", "Técnica / Enfoque", "Herramienta sugerida", "Responsable"]
    d_niveles = [
        ["Pruebas Unitarias", "Caja Blanca: validación de funciones de cálculo de notas, hash de contraseñas, normalización de respuestas escritas y triggers RLS.", "pytest / Jest", "Desarrollador"],
        ["Pruebas de Integración de API", "Caja Negra: validación de contratos de endpoints REST, headers JWT, códigos de estado HTTP y transacciones en BD.", "Postman / Newman / Supertest", "Desarrollador / QA"],
        ["Pruebas de Sistema E2E", "Caja Negra: flujos completos de usuario (Registro -> Aprobación -> Creación de Clase -> Código VF-XXXX -> Examen -> Calificación).", "Playwright / Cypress / Manual", "Desarrollador / Líder Técnico"],
        ["Pruebas de Aceptación UAT", "Pruebas de usuario final en entorno escolar con docentes y estudiantes representativos de grados 10° y 11°.", "Navegadores en sala de cómputo y celulares", "Docentes y Estudiantes"]
    ]
    add_table_data(doc, h_niveles, d_niveles, col_widths=[1.5, 2.5, 1.7, 1.5])

    # 3. Entorno de Pruebas y Datos
    add_heading_with_spacing(doc, "3. Entorno de Pruebas y Criterios de Aceptación", level=2)
    add_bullet(doc, "Dispositivos de Prueba", "Computadores de escritorio (Windows 10/11 en navegadores Chrome y Edge), teléfonos inteligentes Android de gama media/baja (pantallas desde 360 px) y tablets.")
    add_bullet(doc, "Condiciones de Red Simuladas", "Red escolar óptima (Wi-Fi 50 Mbps), red celular 4G estándar y red degradada/inestable con latencia de 1.500 ms y pérdida de paquetes del 10% (Chrome DevTools Network Throttling).")
    add_bullet(doc, "Criterio de Aprobación de Versión (Release Gate)", "Cero defectos críticos (Severidad 1: bloqueo de examen o fuga de respuestas) y al menos 95% de casos de prueba funcionales en estado EXITOSO.")

    # 4. Matriz Exhaustiva de Casos de Prueba Funcionales (IEEE 829)
    add_heading_with_spacing(doc, "4. Matriz Exhaustiva de Casos de Prueba Funcionales (IEEE 829)", level=2)
    add_p(doc, "Casos de prueba detallados clasificados por módulos de la plataforma:")

    # 4.1 Módulo Autenticación y Cuentas
    add_heading_with_spacing(doc, "4.1 Módulo de Autenticación, Registro y Aprobación de Docentes", level=3)
    h_cp = ["ID", "Caso de Prueba", "Precondición", "Pasos y Datos", "Resultado Esperado", "Prior."]
    d_cp_auth = [
        ["CP-AUTH-01", "Registro exitoso de estudiante grado 10° o 11°", "Nickname disponible.", "1. Llenar formulario con nick 'kevin10', contraseña robusta, grado 10 y aceptar Ley 1581.\n2. Enviar.", "Cuenta creada; estado 'activo'; sesión iniciada con JWT; grado 10 asignado; institución fija.", "Alta"],
        ["CP-AUTH-02", "Registro de docente en estado 'pendiente'", "Correo y nick válidos.", "1. Registrarse seleccionando tipo 'Docente'.\n2. Enviar formulario.", "Cuenta creada pero en estado 'pendiente'; no puede crear clases ni evaluaciones hasta aprobación.", "Alta"],
        ["CP-AUTH-03", "Aprobación de docente por el Owner (TheKeas)", "Docente en estado 'pendiente'.", "1. Iniciar sesión como Owner.\n2. Ir a panel /api/v1/auth/docentes/pendientes/.\n3. Pulsar 'Aprobar'.", "Estado del docente cambia a 'activo'; solicitud_docente aprobada; docente ya puede crear clases.", "Alta"],
        ["CP-AUTH-04", "Rechazo de docente por el Owner con motivo", "Docente en estado 'pendiente'.", "1. Iniciar sesión como Owner.\n2. Pulsar 'Rechazar' e ingresar motivo 'No pertenece a la I.E.'.", "Estado cambia a 'rechazado'; docente bloqueado; motivo registrado en base de datos.", "Media"],
        ["CP-AUTH-05", "Validación de unicidad de nickname (case-insensitive)", "Existe usuario 'juanperez'.", "1. Intentar registrar 'JuanPerez' o 'JUANPEREZ'.\n2. Enviar.", "Error HTTP 400: 'NICK_DUPLICADO'. El sistema no distingue mayúsculas/minúsculas.", "Alta"],
        ["CP-AUTH-06", "Obligatoriedad de consentimiento Ley 1581", "Formulario diligenciado.", "1. Intentar registrar desmarcando el checkbox de consentimiento de acudiente.", "Bloqueo en frontend y backend (HTTP 400); no se crea cuenta sin consentimiento legal.", "Alta"]
    ]
    add_table_data(doc, h_cp, d_cp_auth, col_widths=[1.1, 1.8, 1.2, 1.6, 1.8, 0.7])

    # 4.2 Módulo Práctica Oral con Micrófono y Feedback Fonético
    add_heading_with_spacing(doc, "4.2 Módulo de Diálogos Contextuales y Reconocimiento de Voz", level=3)
    d_cp_oral = [
        ["CP-VOZ-01", "Activación del micrófono y animación de ondas sonoras", "Permiso de micrófono habilitado.", "1. Entrar a 'El Alibi del Detective'.\n2. Escuchar turno del detective.\n3. Presionar botón 3D de micrófono.", "Visualizador de ondas de audio pulsa al detectar voz; estado cambia a 'Escuchando tu voz...'.", "Alta"],
        ["CP-VOZ-02", "Detección de pronunciación correcta y avance de turno", "Frase: 'I was watching TV with my family'.", "1. Hablar con claridad la frase completa.\n2. Soltar micrófono.", "Todas las palabras en verde esmeralda; precisión > 85%; botón 'Siguiente Turno' habilitado.", "Alta"],
        ["CP-VOZ-03", "Detección de error fonético específico (watching vs washing)", "Frase esperada con 'watching'.", "1. Pronunciar deliberadamente 'I was washing TV...'.\n2. Soltar micrófono.", "Palabra 'watching' en rojo coral con alerta; Coach Fonético explica sonido /tʃ/ vs /ʃ/ en español.", "Alta"],
        ["CP-VOZ-04", "Reproducción de audio nativo a doble velocidad (1.0x y 0.75x)", "Escenario cargado.", "1. Pulsar botón 'Escuchar 1.0x'.\n2. Pulsar botón 'Audio Lento 0.75x'.", "El navegador reproduce la síntesis/audio con entonación contextual en velocidad normal y lenta.", "Media"],
        ["CP-VOZ-05", "Verificación de privacidad: cero audio enviado a red", "Práctica oral en ejecución.", "1. Abrir DevTools pestaña Network.\n2. Grabar múltiples frases de voz.", "Cero peticiones POST de audio hacia servidores externos. Todo ocurre en memoria local.", "Alta"]
    ]
    add_table_data(doc, h_cp, d_cp_oral, col_widths=[1.1, 1.8, 1.2, 1.6, 1.8, 0.7])

    # 4.3 Módulo Clases y Evaluaciones (Docente)
    add_heading_with_spacing(doc, "4.3 Módulo de Clases y Creación de Evaluaciones Oficiales", level=3)
    d_cp_docente = [
        ["CP-DOC-01", "Creación de grupo escolar (Inglés 10-A)", "Docente activo.", "1. Ingresar a 'Crear Clase'.\n2. Nombre 'Inglés 10-A', grado 10, año 2026.\n3. Guardar.", "Clase creada en BD; aparece en listado 'Mis Clases'; lista para recibir alumnos.", "Alta"],
        ["CP-DOC-02", "Publicación de evaluación con código VF-XXXX", "Clase existente.", "1. Crear evaluación de 15 preguntas, duración 30 min, Past Continuous.\n2. Publicar.", "Código generado (ej. VF-4821); preguntas congeladas en 'evaluacion_pregunta'; estado 'publicada'.", "Alta"],
        ["CP-DOC-03", "Regla 1 Estudiante = 1 Sola Clase", "Estudiante en Inglés 10-A.", "1. Intentar asociar al estudiante a la clase 'Inglés 11-B'.", "El sistema actualiza la clase o advierte cambio de grupo; nunca permite membresía simultánea.", "Alta"],
        ["CP-DOC-04", "Aislamiento de cursos entre docentes", "Docente A y Docente B.", "1. Docente A inicia sesión.\n2. Consultar /api/v1/clases/ y notas.", "Docente A solo ve sus grupos y notas; intento de acceder a clase de Docente B retorna HTTP 403.", "Alta"]
    ]
    add_table_data(doc, h_cp, d_cp_docente, col_widths=[1.1, 1.8, 1.2, 1.6, 1.8, 0.7])

    # 4.4 Módulo Presentación de Examen Oficial y Modo Examen
    add_heading_with_spacing(doc, "4.4 Módulo de Presentación de Examen Oficial y Modo Examen", level=3)
    d_cp_exam = [
        ["CP-EXAM-01", "Acceso válido con código VF-XXXX en horario", "Evaluación publicada activa.", "1. Estudiante digita código VF-XXXX.\n2. Pulsar 'Ingresar'.", "Acceso concedido; muestra título, tiempo disponible y botón 'Iniciar Examen'.", "Alta"],
        ["CP-EXAM-02", "Rechazo de código fuera de ventana temporal", "Prueba abre a las 10:00 AM.", "1. Intentar ingresar a las 9:45 AM con código correcto.", "Error HTTP 423 'EXAMEN_FUERA_DE_HORARIO'; informa que la prueba aún no está abierta.", "Alta"],
        ["CP-EXAM-03", "Activación estricta de Modo Examen", "Examen iniciado.", "1. Verificar interfaz del examen.", "Teoría y Chuleta de Fórmulas bloqueadas; interfaz en pantalla completa; temporizador visible.", "Alta"],
        ["CP-EXAM-04", "Ocultamiento de respuestas en el payload (RNF-04)", "Examen iniciado.", "1. Abrir DevTools Network.\n2. Inspeccionar respuesta JSON de /api/v1/examen/iniciar/.", "Las opciones NO contienen el campo 'es_correcta'; respuestas de completar NO están presentes.", "Alta"],
        ["CP-EXAM-05", "Autoguardado progresivo ante cada respuesta", "Pregunta 1 respondida.", "1. Seleccionar opción B en pregunta 1.\n2. Verificar tráfico de red.", "Petición POST /examen/intentos/{id}/responder/ responde HTTP 200 en segundo plano; avance a salvo.", "Alta"],
        ["CP-EXAM-06", "Cálculo matemático de calificación 1.0 a 5.0", "Examen de 20 preguntas.", "1. Alumno acierta 15 de 20 preguntas.\n2. Finalizar examen.", "Nota calculada: 1.0 + (15/20)*4.0 = 4.00. Guardada en 'intento.puntaje'; estado 'finalizado'.", "Alta"],
        ["CP-EXAM-07", "Control de límite de intentos (intentos_max = 1)", "Examen entregado.", "1. Intentar reingresar con el mismo código VF-XXXX.", "Bloqueo inmediato; HTTP 409 'INTENTOS_AGOTADOS'; muestra mensaje de examen ya presentado.", "Alta"],
        ["CP-EXAM-08", "Cierre automático al expirar el tiempo límite", "Examen con 10 min límite.", "1. Dejar transcurrir los 10 minutos sin entregar.", "El servidor expira el intento automáticamente; califica las respuestas guardadas hasta ese instante.", "Alta"],
        ["CP-EXAM-09", "Registro de señales de foco (salida de pestaña)", "Examen en curso.", "1. Cambiar de pestaña o minimizar ventana.", "El navegador detecta pérdida de foco y envía evento a /examen/intentos/{id}/evento/; docente lo ve.", "Media"]
    ]
    add_table_data(doc, h_cp, d_cp_exam, col_widths=[1.1, 1.8, 1.2, 1.6, 1.8, 0.7])

    # 4.5 Módulo Calificaciones y Reportes
    add_heading_with_spacing(doc, "4.5 Módulo de Monitoreo y Calificaciones", level=3)
    d_cp_rep = [
        ["CP-REP-01", "Tablero de calificaciones en vivo del docente", "Alumnos rindiendo prueba.", "1. Docente ingresa a 'Calificaciones en Vivo'.", "Lista alumnos con estado (en curso / finalizado), tiempo empleado, aciertos y nota preliminar/final.", "Alta"],
        ["CP-REP-02", "Exportación de planilla de notas a Excel / CSV", "Evaluación finalizada.", "1. Pulsar botón 'Descargar Excel/CSV'.", "Descarga archivo formateado con columnas: Nombre, Nickname, Grado, Aciertos, Total y Nota (1.0-5.0).", "Media"]
    ]
    add_table_data(doc, h_cp, d_cp_rep, col_widths=[1.1, 1.8, 1.2, 1.6, 1.8, 0.7])

    # 5. Pruebas No Funcionales (Resiliencia, Seguridad y Rendimiento)
    add_heading_with_spacing(doc, "5. Pruebas No Funcionales (Seguridad, Estrés y Resiliencia)", level=2)
    h_cp_nf = ["ID", "Prueba No Funcional", "Escenario de Prueba", "Criterio de Éxito"]
    d_cp_nf = [
        ["CP-NF-01", "Resiliencia ante corte de red Wi-Fi en el colegio", "Estudiante respondiendo pregunta 8 de 15. Se apaga el adaptador de red por 45 segundos y se reactiva.", "Al reconectar, el frontend reenvía la respuesta en cola sin errores; ninguna pregunta previa se pierde."],
        ["CP-NF-02", "Carga simultánea de 300 estudiantes concurrentes", "Prueba de estrés simulada con 300 peticiones concurrentes de inicio y entrega de examen en 10 segundos.", "Tiempos de respuesta de la API menores a 400 ms; tasa de error 0%; PostgreSQL opera sin bloqueos."],
        ["CP-NF-03", "Inyección SQL y Cross-Site Scripting (XSS)", "Ingreso de payloads maliciosos (' OR 1=1 --, <script>alert(1)</script>) en inputs de búsqueda y respuestas de completar.", "El backend y ORM sanitizan y parametrizan la entrada; cero ejecución de código malicioso."],
        ["CP-NF-04", "Responsividad móvil en pantalla pequeña (360 px)", "Ejecución de la app en emulador móvil de 360 px de ancho.", "Cero desbordamientos horizontales; botones táctiles 3D accesibles; visualizador de ondas y micrófono operables."]
    ]
    add_table_data(doc, h_cp_nf, d_cp_nf, col_widths=[1.2, 2.2, 2.8, 2.0])

    # 6. Matriz de Trazabilidad (SRS v0.3 vs Casos de Prueba)
    add_heading_with_spacing(doc, "6. Matriz de Trazabilidad (SRS v0.3 vs Casos de Prueba)", level=2)
    h_traz = ["Requisito SRS v0.3", "Descripción del Requisito", "Casos de Prueba que lo Cubren"]
    d_traz = [
        ["RF-01 a RF-06", "Autenticación, consentimiento Ley 1581 y aprobación de docentes", "CP-AUTH-01, CP-AUTH-02, CP-AUTH-03, CP-AUTH-04, CP-AUTH-05, CP-AUTH-06"],
        ["RF-15 a RF-19b", "Diálogos contextuales, micrófono, Word-Diff y feedback fonético", "CP-VOZ-01, CP-VOZ-02, CP-VOZ-03, CP-VOZ-04, CP-VOZ-05"],
        ["RF-30 a RF-37", "Exámenes oficiales con código VF-XXXX, autoguardado y nota 1.0 a 5.0", "CP-EXAM-01 al CP-EXAM-09, CP-NF-01"],
        ["RF-40 a RF-47", "Suite docente: clases, evaluaciones, 1 estudiante=1 clase, notas en vivo", "CP-DOC-01, CP-DOC-02, CP-DOC-03, CP-DOC-04, CP-REP-01, CP-REP-02"],
        ["RNF-01 a RNF-08", "Seguridad de respuestas, privacidad de voz, 300 alumnos y móvil 360px", "CP-EXAM-04, CP-VOZ-05, CP-NF-01, CP-NF-02, CP-NF-03, CP-NF-04"]
    ]
    add_table_data(doc, h_traz, d_traz, col_widths=[1.6, 3.4, 2.2])

    # 7. Procedimiento de Gestión de Incidencias
    add_heading_with_spacing(doc, "7. Procedimiento de Gestión de Incidencias y Defectos", level=2)
    add_p(doc, "Toda falla detectada durante las sesiones de prueba se documentará con la siguiente clasificación de severidad:")
    add_bullet(doc, "Severidad 1 (Bloqueante / Crítica)", "Imposibilidad de iniciar o entregar un examen oficial, cálculo erróneo de calificaciones o filtración de respuestas correctas en el navegador. Detiene inmediatamente la liberación.")
    add_bullet(doc, "Severidad 2 (Mayor)", "Falla en el reconocimiento de voz de un escenario específico o error en la descarga de planillas Excel. Debe resolverse en la misma iteración.")
    add_bullet(doc, "Severidad 3 (Menor / Cosmética)", "Desalineaciones visuales leves en pantallas menores a 360 px o textos de consejos fonéticos con errores tipográficos. Se programa para la siguiente actualización.")

    doc.save(filepath)
    print(f"Plan de Pruebas guardado exitosamente en {filepath}")

generate_stp_doc("Docs/VerbFlow — Plan de Pruebas de Software (STP).docx")
