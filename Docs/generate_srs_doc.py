# -*- coding: utf-8 -*-
import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_table_styling(table, col_widths=None, header_bg="1F3A60", alt_bg="F1F5F9"):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            if col_widths and idx < len(col_widths):
                cell.width = Inches(col_widths[idx])
                
    # Header row
    for cell in table.rows[0].cells:
        set_cell_background(cell, header_bg)
        for p in cell.paragraphs:
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            for run in p.runs:
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.font.size = Pt(9.5)
                run.font.name = "Segoe UI"
                
    # Data rows
    for r_idx, row in enumerate(table.rows[1:], start=1):
        bg = alt_bg if r_idx % 2 == 1 else "FFFFFF"
        for cell in row.cells:
            set_cell_background(cell, bg)
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                for run in p.runs:
                    run.font.size = Pt(9)
                    run.font.name = "Segoe UI"

def add_heading_with_spacing(doc, text, level):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    for run in p.runs:
        run.font.name = "Segoe UI"
    return p

def add_p(doc, text, bold_prefix=None, style='Normal'):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r1 = p.add_run(bold_prefix)
        r1.font.bold = True
        r1.font.name = "Segoe UI"
        r1.font.size = Pt(10)
    r2 = p.add_run(text)
    r2.font.name = "Segoe UI"
    r2.font.size = Pt(10)
    return p

def add_bullet(doc, title, text):
    p = doc.add_paragraph(style='List Paragraph')
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    r1 = p.add_run(title + ": ")
    r1.font.bold = True
    r1.font.name = "Segoe UI"
    r1.font.size = Pt(10)
    r2 = p.add_run(text)
    r2.font.name = "Segoe UI"
    r2.font.size = Pt(10)
    return p

def add_table_data(doc, headers, data, col_widths=None):
    table = doc.add_table(rows=len(data) + 1, cols=len(headers))
    table.style = 'Table Grid'
    for c_idx, h in enumerate(headers):
        table.cell(0, c_idx).text = h
    for r_idx, row in enumerate(data, start=1):
        for c_idx, val in enumerate(row):
            table.cell(r_idx, c_idx).text = str(val)
    set_table_styling(table, col_widths=col_widths)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(6)
    return table

def generate_srs(filepath):
    doc = docx.Document()
    
    # Title
    p_title = doc.add_heading("VerbFlow — Especificación de Requisitos de Software (SRS)", level=1)
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(2)
    
    p_meta = doc.add_paragraph()
    p_meta.paragraph_format.space_after = Pt(12)
    r_meta = p_meta.add_run("Versión 0.2 (definitiva) · Autor: Kevin · Basado en IEEE 29148 · Octubre de 2026")
    r_meta.font.italic = True
    r_meta.font.size = Pt(9.5)
    r_meta.font.color.rgb = RGBColor(100, 116, 139)
    
    # 1. Introducción
    add_heading_with_spacing(doc, "1. Introducción", level=2)
    
    add_heading_with_spacing(doc, "1.1 Propósito", level=3)
    add_p(doc, "Este documento define los requisitos funcionales y no funcionales de VerbFlow, una plataforma web privada diseñada para la Institución Educativa Promoción Social (Palermo, Huila, Colombia). Su objetivo es proporcionar un entorno especializado para que los estudiantes de grados décimo (10°) y once (11°) refuercen, practiquen y sean evaluados en el dominio de los tiempos verbales del idioma inglés. Este documento constituye la especificación contractual y técnica para el desarrollo, pruebas y puesta en producción.")
    
    add_heading_with_spacing(doc, "1.2 Alcance", level=3)
    add_p(doc, "La plataforma VerbFlow permitirá:")
    add_bullet(doc, "Estudiar teoría especializada", "Reglas gramaticales, estructuras sintácticas (afirmativa, negativa, interrogativa), líneas de tiempo cronológicas interactivas y reproducción de audio de pronunciación (Web Speech API / archivos pregrabados).")
    add_bullet(doc, "Práctica libre casual y privada", "Sesiones personalizadas de práctica donde el estudiante selecciona temas, dificultad y número de preguntas. Se califica de forma inmediata en el navegador sin almacenar registros en el servidor, garantizando privacidad y rapidez.")
    add_bullet(doc, "Presentación de pruebas oficiales de clase", "Evaluaciones asignadas por los docentes activadas mediante un código alfanumérico único (formato VF-XXXX), con guardado progresivo de respuestas, control de tiempo y señales de supervisión académica (Modo Examen).")
    add_bullet(doc, "Suite de gestión para docentes", "Creación y administración de clases (10° y 11°), diseño de evaluaciones automáticas o manuales, monitoreo de entrega en tiempo real y descarga de calificaciones en escala colombiana (1.0 a 5.0) en formato Excel/CSV.")
    add_bullet(doc, "Panel de administración (Owner)", "Fiscalización de la plataforma, aprobación obligatoria de cuentas docentes antes de permitir su operación, administración del banco curricular global y auditoría inmutable.")
    add_bullet(doc, "Banco de preguntas versátil", "Soporte activo en versión 1.0 para preguntas de opción múltiple (con una sola respuesta correcta) y preguntas de respuesta escrita / completar (con validación de texto normalizado).")
    add_p(doc, "Exclusiones explícitas (Fuera de alcance v1.0): Gamificación basada en XP, niveles, rachas diarias y rankings globales (suprimida por requerimiento institucional para evitar fatiga por métricas y cumplir con minimización de datos personales de menores de edad); materias distintas a inglés; pagos o suscripciones; videollamadas; y calificación mediante inteligencia artificial generativa.")

    add_heading_with_spacing(doc, "1.3 Definiciones y siglas", level=3)
    add_bullet(doc, "Owner", "Administrador general y custodio de la plataforma en la institución.")
    add_bullet(doc, "Docente (Teacher)", "Profesor de inglés de la I.E. Promoción Social autorizado por el Owner para crear cursos y evaluar.")
    add_bullet(doc, "Estudiante", "Alumno de 10° u 11° matriculado en una clase específica.")
    add_bullet(doc, "Modo Examen", "Estado seguro en el cliente web donde se bloquea la teoría y la Chuleta de Fórmulas mientras se rinde una evaluación oficial.")
    add_bullet(doc, "Código de Evaluación", "Clave alfanumérica única en formato VF-XXXX empleada para acceder a la prueba de clase.")
    add_bullet(doc, "Escala MEN", "Sistema de calificación oficial de Colombia de 1.00 a 5.00, con nota mínima aprobatoria institucional de 3.00.")

    add_heading_with_spacing(doc, "1.4 Referencias y marco normativo", level=3)
    add_bullet(doc, "Prototipo de Interfaz (Mockup)", "https://kevin81a.github.io/VerbFlow/")
    add_bullet(doc, "Ley 1581 de 2012 y Decreto 1377 de 2013", "Régimen General de Protección de Datos Personales en Colombia.")
    add_bullet(doc, "Ley 1098 de 2006", "Código de la Infancia y la Adolescencia (protección reforzada y principio de interés superior del menor).")
    add_bullet(doc, "Estándar IEEE 29148:2018", "Systems and software engineering — Life cycle processes — Requirements engineering.")

    # 2. Descripción general
    add_heading_with_spacing(doc, "2. Descripción general", level=2)
    add_heading_with_spacing(doc, "2.1 Perspectiva del producto", level=3)
    add_p(doc, "VerbFlow es una aplicación web de arquitectura modular en tres capas: Interfaz interactiva SPA (Single Page Application), API de servicios RESTful y base de datos relacional PostgreSQL. Se distribuye como una web privada accesible desde computadores, tablets y teléfonos móviles inteligentes a través de navegadores modernos sin requerir instalación nativa.")
    
    add_heading_with_spacing(doc, "2.2 Roles de usuario y permisos", level=3)
    h_roles = ["Rol", "Descripción institucional", "Permisos principales"]
    d_roles = [
        ["Owner", "Administrador general del sistema institucional.", "Aprobar/rechazar docentes, suspender cuentas, gestionar banco global, consultar auditoría y métricas globales."],
        ["Docente", "Profesor de inglés de grado 10° u 11°.", "Crear clases, generar códigos VF-XXXX, diseñar y publicar evaluaciones, agregar preguntas propias (privadas), monitorear en tiempo real y exportar notas."],
        ["Estudiante", "Alumno regular de la I.E. Promoción Social.", "Acceder a estudio, practicar libremente sin registro en servidor, unirse a su clase única, presentar evaluaciones oficiales y ver sus notas."]
    ]
    add_table_data(doc, h_roles, d_roles, col_widths=[1.5, 2.3, 3.2])

    add_heading_with_spacing(doc, "2.3 Entorno operativo", level=3)
    add_p(doc, "Compatible con navegadores web estándar (Chrome, Edge, Firefox, Safari en sus versiones recientes). Optimizado para operar con fluidez en dispositivos móviles de gama media y con alta tolerancia a la latencia o inestabilidad de conexiones en entornos escolares rurales y urbanos de Colombia.")

    add_heading_with_spacing(doc, "2.4 Restricciones del proyecto", level=3)
    add_bullet(doc, "Cumplimiento estricto Ley 1581", "Consentimiento previo expreso e informado. Registro de aceptación con fecha, versión de política e IP.")
    add_bullet(doc, "Minimización de datos de menores", "Por tratarse de estudiantes menores de edad (14-17 años), se prohíbe almacenar documentos de identidad, teléfonos, domicilios o historiales conductuales de gamificación. Solo se tratan datos indispensables para fines estrictamente pedagógicos.")
    add_bullet(doc, "Infraestructura económica", "Diseñado para operar en capas gratuitas o planes básicos de bajo costo (Frontend en hosting estático, Backend en PaaS tipo Render/Railway y PostgreSQL gestionado).")
    add_bullet(doc, "Mantenibilidad para un desarrollador", "Código modular, limpio y estándar en Django y React para facilitar la operación por un desarrollador único.")

    # 3. Requisitos funcionales
    add_heading_with_spacing(doc, "3. Requisitos funcionales", level=2)
    add_p(doc, "Convención de prioridades: M = Mandatorio (imprescindible v1.0), D = Deseable, F = Futuro.")

    # 3.1 Autenticación
    add_heading_with_spacing(doc, "3.1 Autenticación y cuentas de usuario", level=3)
    h_rf = ["ID", "Requisito funcional", "Prioridad"]
    d_rf_auth = [
        ["RF-01", "El sistema permitirá el registro con nickname único, contraseña segura, nombre completo, grado (10° u 11°) y tipo de cuenta (Estudiante o Docente). La institución queda fija como I.E. Promoción Social (no se solicita en formulario).", "M"],
        ["RF-02", "El sistema exigirá marcar la aceptación de la política de tratamiento de datos personales y autorización de acudiente antes de completar el registro (Ley 1581).", "M"],
        ["RF-03", "El sistema permitirá el inicio y cierre seguro de sesión mediante autenticación JWT con tokens de acceso y refresco.", "M"],
        ["RF-04", "Las cuentas registradas como Docente quedarán automáticamente en estado 'pendiente' hasta que el Owner las apruebe.", "M"],
        ["RF-05", "El nickname será único a nivel del sistema y se validará sin distinguir mayúsculas ni minúsculas (case-insensitive).", "M"],
        ["RF-06", "El usuario podrá recuperar su contraseña si suministró un correo electrónico opcional en su perfil.", "D"],
        ["RF-07", "El usuario podrá ejercer su derecho de supresión de datos (habeas data). El sistema anonimizará su información personal preservando únicamente la calificación histórica requerida por la institución.", "M"]
    ]
    add_table_data(doc, h_rf, d_rf_auth, col_widths=[1.2, 5.0, 1.0])

    # 3.2 Modo Estudio
    add_heading_with_spacing(doc, "3.2 Modo Estudio", level=3)
    d_rf_estudio = [
        ["RF-10", "El sistema exhibirá el contenido curricular de cada tiempo verbal: contexto de uso, estructuras (afirmativa, negativa, interrogativa), reglas ortográficas y ejemplos ilustrativos.", "M"],
        ["RF-11", "El sistema incluirá una línea de tiempo visual e interactiva que sitúe cronológicamente la acción verbal en pasado, presente o futuro.", "D"],
        ["RF-12", "El sistema ofrecerá reproducción auditiva de frases y verbos para entrenamiento de pronunciación mediante Web Speech API o archivos MP3.", "M"],
        ["RF-13", "El sistema proveerá una 'Chuleta de Fórmulas' accesible rápidamente desde la interfaz de estudio y práctica.", "M"]
    ]
    add_table_data(doc, h_rf, d_rf_estudio, col_widths=[1.2, 5.0, 1.0])

    # 3.3 Modo Práctica Libre
    add_heading_with_spacing(doc, "3.3 Modo Práctica libre (evaluación casual en cliente)", level=3)
    d_rf_practica = [
        ["RF-20", "El estudiante podrá seleccionar libremente uno, varios o todos los temas verbales para su sesión de práctica.", "M"],
        ["RF-21", "El estudiante podrá elegir el nivel de dificultad: Fácil, Media, Difícil o Ponderada.", "M"],
        ["RF-22", "El estudiante podrá elegir la cantidad de preguntas de la práctica: 5, 10, 15, 20 o 25.", "M"],
        ["RF-23", "El sistema generará la prueba a partir del banco incluyendo tanto preguntas de opción múltiple como de respuesta escrita / completar.", "M"],
        ["RF-24", "Al concluir la práctica, el sistema mostrará inmediatamente el resumen de aciertos, errores y la explicación pedagógica detallada de cada pregunta.", "M"],
        ["RF-25", "La práctica libre operará y se calificará localmente en el navegador del estudiante sin registrar filas de intento en el servidor, garantizando privacidad y cero saturación de la base de datos.", "M"]
    ]
    add_table_data(doc, h_rf, d_rf_practica, col_widths=[1.2, 5.0, 1.0])

    # 3.4 Pruebas de Clase
    add_heading_with_spacing(doc, "3.4 Pruebas de clase (Evaluación oficial formal)", level=3)
    d_rf_examen = [
        ["RF-30", "El estudiante ingresará a la prueba oficial suministrando el código alfanumérico generado por su docente (formato VF-XXXX).", "M"],
        ["RF-31", "El sistema validará que el código exista, esté activo, dentro de la ventana de apertura/cierre y que el estudiante pertenezca a la clase destinataria.", "M"],
        ["RF-32", "Al iniciar la prueba oficial, se activará el Modo Examen en el navegador: se bloqueará la navegación hacia la teoría y la Chuleta de Fórmulas.", "M"],
        ["RF-33", "El sistema controlará el número de intentos permitidos (definido por el docente, por defecto 1). No permitirá iniciar intentos adicionales si se alcanzó el tope.", "M"],
        ["RF-34", "El sistema guardará progresivamente cada respuesta en el servidor al ser marcada o digitada, protegiendo al alumno ante caídas de conexión eléctrica o de internet.", "M"],
        ["RF-35", "El sistema podrá fijar un temporizador regresivo estricto gestionado por el servidor; al cumplirse el tiempo límite, cerrará y calificará automáticamente la prueba.", "M"],
        ["RF-36", "El sistema registrará telemetría de incidencias (pérdida de foco, cambio de pestaña, salida de pantalla completa) como señales de advertencia para el docente.", "D"],
        ["RF-37", "Los resultados se calificarán en escala 1.0 a 5.0 y se entregarán de forma inmediata o diferida según la configuración fijada por el docente.", "M"]
    ]
    add_table_data(doc, h_rf, d_rf_examen, col_widths=[1.2, 5.0, 1.0])

    # 3.5 Suite del Docente
    add_heading_with_spacing(doc, "3.5 Suite del Docente", level=3)
    d_rf_docente = [
        ["RF-40", "El docente podrá crear, editar y archivar clases de grado 10° y 11°. Cada estudiante pertenecerá a una única clase a la vez.", "M"],
        ["RF-41", "El docente podrá crear evaluaciones y generar su código de acceso único VF-XXXX.", "M"],
        ["RF-42", "El docente podrá configurar la prueba en modo automático (por temas y dificultad) o manual (eligiendo preguntas específicas del banco).", "M"],
        ["RF-43", "El docente definirá ventana de calendario (apertura y cierre), duración en minutos, número de intentos y modo de entrega de resultados.", "M"],
        ["RF-44", "El docente dispondrá de un panel de monitoreo en tiempo real donde visualizará quién ha entregado, quién está en curso, nota y alertas de foco.", "M"],
        ["RF-45", "El docente podrá inspeccionar el desglose de respuestas por estudiante y las preguntas que presentaron mayor índice de error en el grupo.", "D"],
        ["RF-46", "El docente podrá descargar la planilla de calificaciones completa en formato Excel / CSV para el registro académico institucional.", "D"],
        ["RF-47", "El docente podrá redactar e incorporar preguntas propias (de opción múltiple o completar) al banco; estas permanecerán privadas para sus clases salvo publicación por el Owner.", "M"]
    ]
    add_table_data(doc, h_rf, d_rf_docente, col_widths=[1.2, 5.0, 1.0])

    # 3.6 Panel Owner
    add_heading_with_spacing(doc, "3.6 Panel del Owner (Administrador)", level=3)
    d_rf_owner = [
        ["RF-50", "El Owner visualizará las solicitudes de docentes en estado 'pendiente' con un contador destacado.", "M"],
        ["RF-51", "El Owner podrá aprobar o rechazar solicitudes de docentes mediante un solo clic.", "M"],
        ["RF-52", "El Owner podrá suspender o reactivar cuentas de usuarios en caso de faltas disciplinarias o requerimiento directivo.", "M"],
        ["RF-53", "El Owner tendrá control total sobre el banco global de preguntas y el temario pedagógico, con capacidad de publicar preguntas docentes a nivel global.", "M"]
    ]
    add_table_data(doc, h_rf, d_rf_owner, col_widths=[1.2, 5.0, 1.0])

    # 3.7 Resolución de Gamificación
    add_heading_with_spacing(doc, "3.7 Decisión sobre Gamificación (RF-60 a RF-65)", level=3)
    add_p(doc, "Por decisión concertada de diseño pedagógico e institucional, se excluyen del alcance del sistema los requisitos RF-60 a RF-65 (puntos XP, niveles calculados, rachas de días consecutivos, metas de estudio y rankings competitivos).")
    add_p(doc, "Justificación técnica e institucional:", bold_prefix="Motivos de exclusión: ")
    add_bullet(doc, "Prevención de saturación y fatiga", "VerbFlow está concebido como una herramienta de apoyo académico formal y práctica casual; los esquemas de racha y rankings generan ansiedad innecesaria o desinterés en el entorno escolar.")
    add_bullet(doc, "Protección reforzada de menores (Ley 1581 / Ley 1098)", "Evitar almacenar métricas conductuales y de hábitos de estudio reduce el perfilamiento de menores y simplifica la política de privacidad institucional.")
    add_bullet(doc, "Simplicidad y desempeño", "Se eliminan 4 tablas de la base de datos y tareas programadas nocturnas de cálculo de rachas, permitiendo que la plataforma corra con máxima agilidad y bajo consumo de recursos.")

    # 4. Requisitos no funcionales
    add_heading_with_spacing(doc, "4. Requisitos no funcionales", level=2)
    h_rnf = ["ID", "Categoría", "Especificación técnica"]
    d_rnf = [
        ["RNF-01", "Seguridad", "Almacenamiento de contraseñas mediante funciones criptográficas robustas con salting (Argon2 o PBKDF2). Prohibido texto plano."],
        ["RNF-02", "Seguridad", "Toda comunicación cliente-servidor se ejecutará obligatoriamente sobre protocolo cifrado HTTPS."],
        ["RNF-03", "Seguridad", "Control estricto de acceso basado en roles (RBAC) validado en cada endpoint del backend."],
        ["RNF-04", "Seguridad", "Durante evaluaciones oficiales, las respuestas correctas nunca se envían al navegador; la verificación se efectúa exclusivamente en el servidor."],
        ["RNF-05", "Seguridad", "Protección activa contra inyección SQL, Cross-Site Scripting (XSS), Cross-Site Request Forgery (CSRF) y limitación de tasa contra ataques de fuerza bruta."],
        ["RNF-06", "Privacidad", "Cumplimiento integral de la Ley Estatutaria 1581 de 2012: consentimiento previo, finalidad académica legítima y principio de necesidad."],
        ["RNF-07", "Rendimiento", "Tiempo de carga inicial inferior a 3 segundos en redes 4G móviles. Respuestas de API inferiores a 400 ms en el 95% de peticiones."],
        ["RNF-08", "Escalabilidad", "Capacidad para soportar al menos 300 estudiantes concurrentes rindiendo una prueba simultánea sin degradación del servicio."],
        ["RNF-09", "Disponibilidad", "Disponibilidad del servicio de al menos 99% en jornada académica escolar (lunes a viernes de 6:30 a 18:30 hora de Colombia)."],
        ["RNF-10", "Usabilidad", "Interfaz 100% en español, responsiva, con alto contraste visual y cumplimiento de pautas básicas de accesibilidad web (WCAG 2.1 AA)."],
        ["RNF-11", "Compatibilidad", "Totalmente operativa en pantallas móviles desde 360 px de ancho."],
        ["RNF-12", "Mantenibilidad", "Repositorio Git estructurado, separación de capas clara, suite de pruebas unitarias automáticas y documentación interactiva OpenAPI (Swagger)."],
        ["RNF-13", "Respaldo", "Copias de seguridad automáticas diarias de la base de datos PostgreSQL cifradas en reposo."],
        ["RNF-14", "Integridad", "Guardado progresivo de respuestas y transacciones atómicas para impedir pérdida de calificaciones ante cortes eléctricos o de red en el colegio."]
    ]
    add_table_data(doc, h_rnf, d_rnf, col_widths=[1.2, 1.6, 4.4])

    # 5. Reglas de negocio
    add_heading_with_spacing(doc, "5. Reglas de negocio", level=2)
    add_bullet(doc, "RN-01", "Un docente no puede acceder a las funciones de creación de clases o pruebas hasta que su cuenta sea aprobada formalmente por el Owner.")
    add_bullet(doc, "RN-02", "Un estudiante pertenece exclusivamente a la I.E. Promoción Social y solo puede estar matriculado en exactamente una clase a la vez.")
    add_bullet(doc, "RN-03", "El código de evaluación VF-XXXX solo habilita el inicio del examen dentro de la ventana de apertura y cierre configurada por el docente.")
    add_bullet(doc, "RN-04", "Durante el Modo Examen queda inhabilitado el acceso a las vistas de teoría, línea de tiempo y Chuleta de Fórmulas.")
    add_bullet(doc, "RN-05", "Un estudiante no puede exceder el límite de intentos fijado por el docente para una evaluación oficial.")
    add_bullet(doc, "RN-06", "Los datos de menores de edad solo se tratan con la previa manifestación de autorización del acudiente o mayor de edad legalmente facultado.")
    add_bullet(doc, "RN-07", "Aislamiento docente: Un docente solo puede consultar notas y estudiantes de sus propios cursos asignados.")
    add_bullet(doc, "RN-08", "Escala de notas: Las calificaciones se computan en la escala oficial de 1.00 a 5.00 mediante la fórmula: Puntaje = 1.0 + (aciertos / total_preguntas) * 4.0. La nota mínima de aprobación institucional es 3.00.")

    # 6. Casos de uso principales
    add_heading_with_spacing(doc, "6. Casos de uso principales", level=2)
    h_cu = ["ID", "Caso de uso", "Actor principal", "Resultado esperado"]
    d_cu = [
        ["CU-01", "Registrarse en la plataforma", "Estudiante, Docente", "Cuenta registrada (Docente queda en estado 'pendiente')."],
        ["CU-02", "Aprobar docente", "Owner", "Cuenta docente activada con permisos para operar cursos."],
        ["CU-03", "Estudiar tiempo verbal", "Estudiante", "Acceso a reglas, audio de pronunciación, ejemplos y chuleta."],
        ["CU-04", "Practicar libremente", "Estudiante", "Sesión interactiva calificada de forma inmediata en el navegador sin persistencia."],
        ["CU-05", "Crear clase y evaluación", "Docente", "Generación del código alfanumérico VF-XXXX y congelación de preguntas."],
        ["CU-06", "Presentar evaluación de clase", "Estudiante", "Respuestas guardadas progresivamente y nota en escala 1.0 a 5.0 calculada en backend."],
        ["CU-07", "Monitorear examen en vivo", "Docente", "Tablero de avance con notas, tiempos y señales de foco en tiempo real."],
        ["CU-08", "Exportar notas académicas", "Docente", "Generación y descarga de archivo Excel/CSV con calificaciones del curso."]
    ]
    add_table_data(doc, h_cu, d_cu, col_widths=[1.1, 2.2, 1.7, 2.2])

    # 7. Arquitectura y stack propuesto
    add_heading_with_spacing(doc, "7. Arquitectura y stack tecnológico", level=2)
    add_bullet(doc, "Frontend", "React con Vite, React Router y TanStack Query para gestión de estado del servidor. Estilos con Tailwind CSS siguiendo fielmente el mockup de referencia.")
    add_bullet(doc, "Backend", "Python con Django y Django REST Framework (DRF). Manejo nativo de seguridad, autenticación JWT, ORM con transacciones seguras y panel administrativo de apoyo.")
    add_bullet(doc, "Base de datos", "PostgreSQL con tipos de datos nativos (UUID, JSONB, TIMESTAMPTZ, INET) y restricciones de integridad relacional.")
    add_bullet(doc, "Despliegue", "Frontend alojado en Vercel/Netlify; Backend y PostgreSQL alojados en Render/Railway con copias de respaldo automáticas.")

    # 8. Modelo de datos preliminar
    add_heading_with_spacing(doc, "8. Modelo de datos consolidado (17 entidades)", level=2)
    add_p(doc, "El sistema se estructura en 17 entidades relacionales:")
    add_bullet(doc, "Entidades principales", "usuario (con clase_id directo), consentimiento, solicitud_docente, auditoria, tema, formula, audio_ejemplo, pregunta, opcion_pregunta, respuesta_aceptada, clase, evaluacion, evaluacion_tema, evaluacion_pregunta, intento, intento_respuesta, evento_examen.")

    # 9. Criterios de aceptación
    add_heading_with_spacing(doc, "9. Criterios de aceptación de la versión 1.0", level=2)
    add_bullet(doc, "Flujo de estudiante", "Un estudiante puede registrarse, estudiar la teoría, practicar de forma privada y responder un examen oficial de punta a punta.")
    add_bullet(doc, "Flujo de docente", "Un docente aprobado puede crear su clase, publicar un examen, distribuir el código VF-XXXX y monitorear/exportar notas en escala 1.0 a 5.0.")
    add_bullet(doc, "Integridad de evaluación", "Las respuestas correctas son invisibles en el cliente durante exámenes y se guardan progresivamente ante cortes de red.")
    add_bullet(doc, "Legalidad y privacidad", "Cumplimiento estricto de la Ley 1581 de 2012 y minimización absoluta de datos de menores.")

    # 10. Resoluciones de preguntas abiertas
    add_heading_with_spacing(doc, "10. Cierre y resolución de decisiones de diseño", level=2)
    h_res = ["Tema evaluado", "Decisión acordada", "Impacto en el sistema"]
    d_res = [
        ["Cantidad de docentes", "Múltiples docentes dentro de la I.E. Promoción Social.", "Aislamiento estricto de notas por docente y aprobación obligatoria del Owner."],
        ["Tipos de preguntas", "Tanto opción múltiple como escribir/completar en v1.0.", "Activación inmediata de 'opcion_pregunta' y 'respuesta_aceptada'."],
        ["Institución", "Exclusiva para la I.E. Promoción Social (Palermo, Huila).", "Eliminación de la tabla 'institucion'; configuración fija global."],
        ["Relación estudiante-clase", "1 estudiante = 1 sola clase a la vez.", "Eliminación de la tabla 'inscripcion'; 'clase_id' directo en 'usuario'."],
        ["Gamificación (XP/Rachas)", "Suprimida en su totalidad.", "Eliminadas 4 tablas; práctica libre casual en cliente sin persistencia."],
        ["Escala de calificación", "Escala colombiana 1.0 a 5.0 (aprobación 3.0).", "Cálculo matemático exacto integrado en el backend sin depender de gamificación."]
    ]
    add_table_data(doc, h_res, d_res, col_widths=[1.8, 2.5, 2.9])

    doc.save(filepath)
    print(f"SRS guardado exitosamente en {filepath}")

generate_srs("Docs/VerbFlow — Especificación de Requisitos de Software (SRS).docx")
