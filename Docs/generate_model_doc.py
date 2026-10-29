# -*- coding: utf-8 -*-
import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_table_styling(table, col_widths=None, header_bg="1F3A60", alt_bg="F1F5F9"):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            if col_widths and idx < len(col_widths):
                cell.width = Inches(col_widths[idx])
                
    # Header row
    for cell in table.rows[0].cells:
        set_cell_background(cell, header_bg)
        for p in cell.paragraphs:
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            for run in p.runs:
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.font.size = Pt(9.5)
                run.font.name = "Segoe UI"
                
    # Data rows
    for r_idx, row in enumerate(table.rows[1:], start=1):
        bg = alt_bg if r_idx % 2 == 1 else "FFFFFF"
        for cell in row.cells:
            set_cell_background(cell, bg)
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                for run in p.runs:
                    run.font.size = Pt(9)
                    run.font.name = "Segoe UI"

def add_heading_with_spacing(doc, text, level):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    for run in p.runs:
        run.font.name = "Segoe UI"
    return p

def add_p(doc, text, bold_prefix=None, style='Normal'):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r1 = p.add_run(bold_prefix)
        r1.font.bold = True
        r1.font.name = "Segoe UI"
        r1.font.size = Pt(10)
    r2 = p.add_run(text)
    r2.font.name = "Segoe UI"
    r2.font.size = Pt(10)
    return p

def add_bullet(doc, title, text):
    p = doc.add_paragraph(style='List Paragraph')
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    r1 = p.add_run(title + ": ")
    r1.font.bold = True
    r1.font.name = "Segoe UI"
    r1.font.size = Pt(10)
    r2 = p.add_run(text)
    r2.font.name = "Segoe UI"
    r2.font.size = Pt(10)
    return p

def add_table_data(doc, headers, data, col_widths=None):
    table = doc.add_table(rows=len(data) + 1, cols=len(headers))
    table.style = 'Table Grid'
    # Headers
    for c_idx, h in enumerate(headers):
        table.cell(0, c_idx).text = h
    # Rows
    for r_idx, row in enumerate(data, start=1):
        for c_idx, val in enumerate(row):
            table.cell(r_idx, c_idx).text = str(val)
    set_table_styling(table, col_widths=col_widths)
    # Add small space after table
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(6)
    return table

# ==============================================================================
# 1. BUILD MODELO DE DATOS V0.2
# ==============================================================================
def generate_modelo_datos(filepath):
    doc = docx.Document()
    
    # Title
    p_title = doc.add_heading("VerbFlow — Modelo de Datos", level=1)
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(2)
    
    p_meta = doc.add_paragraph()
    p_meta.paragraph_format.space_after = Pt(12)
    r_meta = p_meta.add_run("Versión 0.2 (definitiva) · Autor: Kevin · Complementa el SRS v0.2 · Motor: PostgreSQL · Octubre de 2026")
    r_meta.font.italic = True
    r_meta.font.size = Pt(9.5)
    r_meta.font.color.rgb = RGBColor(100, 116, 139)
    
    # 1. Alcance y decisiones confirmadas
    add_heading_with_spacing(doc, "1. Alcance y decisiones confirmadas", level=2)
    add_p(doc, "Este documento define las entidades, atributos, relaciones, restricciones e índices de la base de datos de VerbFlow, adaptado y simplificado a partir de las decisiones tomadas para la Institución Educativa Promoción Social (Palermo, Huila, Colombia). Cada tabla se traza a los requisitos funcionales del SRS.")
    add_p(doc, "Decisiones de diseño consolidadas:", bold_prefix="Cierre de supuestos: ")
    
    add_bullet(doc, "1. Tipos de pregunta", "La versión inicial (v1.0) implementa tanto opción múltiple (3 a 5 alternativas con única opción correcta) como respuesta escrita / completar (cotejo normalizado contra respuestas aceptadas).")
    add_bullet(doc, "2. Institución única y privada", "La plataforma es de uso exclusivo para la Institución Educativa Promoción Social (Palermo, Huila). Se elimina la tabla 'institucion'; la parametrización institucional queda fija en la configuración global.")
    add_bullet(doc, "3. Múltiples docentes con aprobación", "Varios docentes de inglés pueden operar la plataforma. Cada docente crea y administra sus propias clases y evaluaciones, y solo accede a las notas de sus estudiantes. Toda cuenta docente nace en estado 'pendiente' y requiere aprobación obligatoria del Owner.")
    add_bullet(doc, "4. Banco de preguntas co-administrado", "El Owner aporta preguntas de visibilidad global y los docentes pueden registrar preguntas propias. Las preguntas creadas por docentes nacen con visibilidad 'privada' (uso exclusivo de su autor) y el Owner puede promoverlas a 'global'.")
    add_bullet(doc, "5. Relación 1 Estudiante = 1 Sola Clase", "Un estudiante solo puede pertenecer a una clase a la vez. Se elimina la tabla intermedia 'inscripcion' y se reemplaza por una relación foránea directa 'clase_id' en la entidad 'usuario'. Al ingresar por primera vez con el código VF-XXXX de una clase, el estudiante queda vinculado a ella.")
    add_bullet(doc, "6. Supresión total de gamificación", "Se eliminan las tablas 'progreso', 'registro_xp', 'estudio_tema' e 'intento_tema'. La plataforma evita saturar al estudiante con métricas de racha o niveles en un entorno escolar casual. Asimismo, se maximiza el cumplimiento de la Ley 1581 de 2012 y la Ley 1098 de 2006 al no almacenar datos conductuales innecesarios de menores.")
    add_bullet(doc, "7. Práctica libre sin persistencia en servidor", "El modo práctica opera de forma inmediata en el navegador del estudiante. El cliente solicita un lote de preguntas y califica localmente en memoria. Solo las evaluaciones oficiales de clase se persisten en base de datos.")
    add_bullet(doc, "8. Calificación en escala colombiana (1.0 a 5.0)", "Las evaluaciones oficiales se califican en la escala estándar del Ministerio de Educación Nacional de Colombia (1.00 a 5.00), donde la nota mínima aprobatoria institucional es 3.00.")

    # 2. Convenciones
    add_heading_with_spacing(doc, "2. Convenciones", level=2)
    add_bullet(doc, "Nomenclatura", "Nombres de tablas y columnas en snake_case, en español y sin tildes.")
    add_bullet(doc, "Claves primarias", "Uso de UUID v4 para entidades expuestas o manipuladas en la API (usuario, clase, evaluacion, intento, pregunta) para prevenir ataques de enumeración secuencial. Tablas internas o de catálogo emplean enteros (smallint o bigint).")
    add_bullet(doc, "Trazabilidad temporal", "Todas las tablas transaccionales cuentan con 'creado_en' y 'actualizado_en' (timestamptz en UTC).")
    add_bullet(doc, "Borrado lógico e inmutabilidad", "Uso de campos de estado o banderas lógicas. Las preguntas utilizadas en evaluaciones nunca se eliminan ni sobreescriben; al ser modificadas se genera una versión nueva (version + 1) archivando la previa, protegiendo las notas históricas.")
    add_bullet(doc, "Enums y validaciones", "Implementados a nivel de aplicación (Django TextChoices) y validados en base de datos mediante restricciones CHECK.")

    # 3. Resumen de entidades
    add_heading_with_spacing(doc, "3. Resumen de entidades (17 tablas)", level=2)
    add_p(doc, "El esquema consolidado pasa de 22 a 17 tablas activas, organizadas en cuatro dominios funcionales:")
    
    headers_ent = ["Dominio funcional", "Tablas", "Propósito"]
    data_ent = [
        ["Usuarios y cumplimiento legal", "usuario, consentimiento, solicitud_docente, auditoria", "Control de accesos, roles, trazabilidad y cumplimiento estricto de la Ley 1581."],
        ["Contenido pedagógico y banco", "tema, formula, audio_ejemplo, pregunta, opcion_pregunta, respuesta_aceptada", "Base curricular de tiempos verbales y banco de preguntas (opción múltiple y completar)."],
        ["Gestión de clases y exámenes", "clase, evaluacion, evaluacion_tema, evaluacion_pregunta", "Estructura académica de grados 10 y 11, parametrización de pruebas y congelación de preguntas."],
        ["Intentos y auditoría de examen", "intento, intento_respuesta, evento_examen", "Almacenamiento exclusivo de evaluaciones oficiales de clase, respuestas y señales antifraude."]
    ]
    add_table_data(doc, headers_ent, data_ent, col_widths=[2.0, 2.5, 2.5])

    # 4. Diccionario de datos
    add_heading_with_spacing(doc, "4. Diccionario de datos", level=2)
    
    # 4.1 usuario
    add_heading_with_spacing(doc, "4.1 usuario (RF-01 a RF-07)", level=3)
    add_p(doc, "Almacena las credenciales y perfiles de acceso. El estudiante tiene vínculo foráneo directo a su clase actual.")
    h_cols = ["Columna", "Tipo", "Restricciones", "Descripción"]
    d_usuario = [
        ["id", "UUID", "PK", "Identificador único."],
        ["nickname", "varchar(30)", "NOT NULL, UNIQUE (lower)", "Nombre de usuario único case-insensitive."],
        ["contrasena_hash", "varchar(255)", "NOT NULL", "Hash criptográfico seguro (Argon2 / PBKDF2)."],
        ["nombre_completo", "varchar(150)", "NOT NULL", "Nombre y apellidos según documento."],
        ["rol", "enum", "NOT NULL: owner, docente, estudiante", "Rol en la plataforma."],
        ["estado", "enum", "NOT NULL: pendiente, activo, suspendido, rechazado, eliminado", "Docentes nacen en 'pendiente'; estudiantes nacen en 'activo'."],
        ["clase_id", "UUID", "FK → clase, NULL", "Clase a la que pertenece el estudiante. NULL si es docente/owner o sin clase."],
        ["grado", "smallint", "CHECK (grado IN (10, 11)), NULL", "Grado cursado. Obligatorio para estudiantes; NULL para docentes/owner."],
        ["correo", "varchar(254)", "NULL, UNIQUE si existe", "Correo opcional; empleado para restablecimiento de contraseña."],
        ["ultimo_acceso_en", "timestamptz", "NULL", "Fecha y hora del último inicio de sesión."],
        ["creado_en", "timestamptz", "NOT NULL, default now()", "Fecha de registro."],
        ["actualizado_en", "timestamptz", "NOT NULL, default now()", "Fecha de última modificación."],
        ["eliminado_en", "timestamptz", "NULL", "Fecha de revocación o baja de cuenta (Ley 1581)."]
    ]
    add_table_data(doc, h_cols, d_usuario, col_widths=[1.5, 1.2, 1.8, 2.5])

    # 4.2 consentimiento
    add_heading_with_spacing(doc, "4.2 consentimiento (RF-02, RNF-06)", level=3)
    add_p(doc, "Registro de aceptación de la política de tratamiento de datos personales para respaldo ante la Ley 1581 de 2012 y Ley 1098 de 2006.")
    d_consentimiento = [
        ["id", "UUID", "PK", "Identificador."],
        ["usuario_id", "UUID", "FK → usuario, NOT NULL", "Usuario que otorga el consentimiento."],
        ["tipo", "enum", "NOT NULL: mayor_edad, acudiente", "Condición declarada."],
        ["version_politica", "varchar(20)", "NOT NULL", "Versión del documento legal aceptado."],
        ["aceptado_en", "timestamptz", "NOT NULL, default now()", "Marca de tiempo legal."],
        ["ip", "inet", "NULL", "Dirección IP para trazabilidad y auditoría."],
        ["revocado_en", "timestamptz", "NULL", "Fecha en caso de solicitar revocatoria."]
    ]
    add_table_data(doc, h_cols, d_consentimiento, col_widths=[1.5, 1.2, 1.8, 2.5])

    # 4.3 solicitud_docente
    add_heading_with_spacing(doc, "4.3 solicitud_docente (RF-04, RF-50, RF-51)", level=3)
    add_p(doc, "Bandeja de verificación de docentes por parte del Owner para impedir suplantaciones.")
    d_solicitud = [
        ["id", "UUID", "PK", "Identificador."],
        ["usuario_id", "UUID", "FK → usuario, NOT NULL, UNIQUE", "Docente que solicita la habilitación."],
        ["estado", "enum", "NOT NULL: pendiente, aprobada, rechazada", "Estado de la solicitud."],
        ["revisado_por", "UUID", "FK → usuario, NULL", "Owner que revisó la solicitud."],
        ["revisado_en", "timestamptz", "NULL", "Fecha de resolución."],
        ["motivo_rechazo", "text", "NULL", "Explicación en caso de rechazo."]
    ]
    add_table_data(doc, h_cols, d_solicitud, col_widths=[1.5, 1.2, 1.8, 2.5])

    # 4.4 auditoria
    add_heading_with_spacing(doc, "4.4 auditoria (Seguridad)", level=3)
    add_p(doc, "Bitácora inmutable de acciones administrativas y de control sensible.")
    d_audit = [
        ["id", "bigint", "PK autoincremental", "Identificador secuencial."],
        ["usuario_id", "UUID", "FK → usuario, NULL", "Usuario que ejecutó la operación."],
        ["accion", "varchar(60)", "NOT NULL", "Ej.: aprobar_docente, suspender_cuenta, publicar_evaluacion."],
        ["entidad", "varchar(60)", "NOT NULL", "Nombre de la tabla afectada."],
        ["entidad_id", "varchar(64)", "NOT NULL", "ID del registro afectado."],
        ["ip", "inet", "NULL", "Dirección IP de origen."],
        ["creado_en", "timestamptz", "NOT NULL, default now()", "Fecha y hora del evento."]
    ]
    add_table_data(doc, h_cols, d_audit, col_widths=[1.5, 1.2, 1.8, 2.5])

    # 4.5 tema
    add_heading_with_spacing(doc, "4.5 tema (RF-10, RF-11)", level=3)
    add_p(doc, "Catálogo pedagógico de tiempos verbales (Present Simple, Past Continuous, etc.).")
    d_tema = [
        ["id", "smallint", "PK", "Identificador numérico."],
        ["slug", "varchar(60)", "NOT NULL, UNIQUE", "Identificador URL (ej.: past-continuous)."],
        ["nombre", "varchar(100)", "NOT NULL", "Nombre formal en español e inglés."],
        ["orden", "smallint", "NOT NULL", "Secuencia pedagógica de presentación."],
        ["uso", "text", "NOT NULL", "Explicación contextual de cuándo se utiliza."],
        ["reglas_ortografia", "text", "NULL", "Normas de terminaciones (-ed, -ing, irregulares)."],
        ["linea_tiempo", "jsonb", "NULL", "Estructura de coordenadas para renderizar la línea de tiempo interactiva."],
        ["activo", "boolean", "NOT NULL, default true", "Disponibilidad pedagógica."]
    ]
    add_table_data(doc, h_cols, d_tema, col_widths=[1.5, 1.2, 1.8, 2.5])

    # 4.6 formula
    add_heading_with_spacing(doc, "4.6 formula (RF-13)", level=3)
    add_p(doc, "Estructuras sintácticas para la Chuleta de Fórmulas interactiva.")
    d_formula = [
        ["id", "integer", "PK autoincremental", "Identificador."],
        ["tema_id", "smallint", "FK → tema, NOT NULL", "Tema verbal asociado."],
        ["tipo", "enum", "NOT NULL: afirmativa, negativa, interrogativa", "Modalidad de la oración."],
        ["estructura", "varchar(200)", "NOT NULL", "Sintaxis (ej.: Sujeto + was/were + verbo-ing)."],
        ["ejemplo", "varchar(250)", "NOT NULL", "Oración modelo representativa."]
    ]
    add_table_data(doc, h_cols, d_formula, col_widths=[1.5, 1.2, 1.8, 2.5])

    # 4.7 audio_ejemplo
    add_heading_with_spacing(doc, "4.7 audio_ejemplo (RF-12)", level=3)
    add_p(doc, "Recursos de pronunciación para entrenamiento auditivo.")
    d_audio = [
        ["id", "integer", "PK autoincremental", "Identificador."],
        ["tema_id", "smallint", "FK → tema, NOT NULL", "Tema verbal asociado."],
        ["texto", "varchar(250)", "NOT NULL", "Frase en inglés a pronunciar."],
        ["archivo_url", "varchar(300)", "NULL", "Audio MP3 pregrabado. Si es NULL, el cliente usa Web Speech API."]
    ]
    add_table_data(doc, h_cols, d_audio, col_widths=[1.5, 1.2, 1.8, 2.5])

    # 4.8 pregunta
    add_heading_with_spacing(doc, "4.8 pregunta (RF-23, RF-47, RF-53)", level=3)
    add_p(doc, "Banco central de preguntas. Soporta opción múltiple y respuesta escrita. Controla versionado para asegurar notas históricas inmutables.")
    d_pregunta = [
        ["id", "UUID", "PK", "Identificador."],
        ["tema_id", "smallint", "FK → tema, NOT NULL", "Tiempo verbal evaluado."],
        ["tipo", "enum", "NOT NULL: opcion_multiple, completar", "Modalidad de pregunta activas en v1.0."],
        ["enunciado", "text", "NOT NULL", "Pregunta o texto con espacio para completar."],
        ["dificultad", "enum", "NOT NULL: facil, media, dificil", "Nivel de complejidad."],
        ["explicacion", "text", "NOT NULL", "Retroalimentación pedagógica mostrada en la corrección."],
        ["autor_id", "UUID", "FK → usuario, NOT NULL", "Docente creador u Owner."],
        ["visibilidad", "enum", "NOT NULL: global, privada", "Privada para el autor docente; global si es aprobada o de Owner."],
        ["estado", "enum", "NOT NULL: activa, archivada", "Borrado lógico."],
        ["version", "integer", "NOT NULL, default 1", "Versión incremental."],
        ["pregunta_origen_id", "UUID", "FK → pregunta, NULL", "Referencia a la versión anterior en caso de edición."]
    ]
    add_table_data(doc, h_cols, d_pregunta, col_widths=[1.5, 1.2, 1.8, 2.5])

    # 4.9 opcion_pregunta
    add_heading_with_spacing(doc, "4.9 opcion_pregunta", level=3)
    add_p(doc, "Opciones de respuesta para preguntas de tipo 'opcion_multiple'.")
    d_opcion = [
        ["id", "UUID", "PK", "Identificador."],
        ["pregunta_id", "UUID", "FK → pregunta, NOT NULL", "Pregunta a la que pertenece."],
        ["texto", "varchar(250)", "NOT NULL", "Texto de la opción."],
        ["es_correcta", "boolean", "NOT NULL", "Indica la respuesta correcta (nunca enviada al cliente en exámenes)."],
        ["orden", "smallint", "NOT NULL", "Orden canónico de presentación."]
    ]
    add_table_data(doc, h_cols, d_opcion, col_widths=[1.5, 1.2, 1.8, 2.5])

    # 4.10 respuesta_aceptada
    add_heading_with_spacing(doc, "4.10 respuesta_aceptada", level=3)
    add_p(doc, "Textos correctos para preguntas de tipo 'completar' (respuesta escrita). El servidor normaliza (minúsculas, trim y sin signos redundantes) antes de validar.")
    d_resp_acep = [
        ["id", "UUID", "PK", "Identificador."],
        ["pregunta_id", "UUID", "FK → pregunta, NOT NULL", "Pregunta de completar asociada."],
        ["texto", "varchar(200)", "NOT NULL", "Palabra o frase exacta aceptada como válida."]
    ]
    add_table_data(doc, h_cols, d_resp_acep, col_widths=[1.5, 1.2, 1.8, 2.5])

    # 4.11 clase
    add_heading_with_spacing(doc, "4.11 clase (RF-40)", level=3)
    add_p(doc, "Grupo académico de grado décimo u once gestionado por un docente específico.")
    d_clase = [
        ["id", "UUID", "PK", "Identificador."],
        ["docente_id", "UUID", "FK → usuario, NOT NULL", "Docente titular a cargo."],
        ["nombre", "varchar(100)", "NOT NULL", "Nombre del grupo (ej.: Inglés 10-A, Inglés 11-B)."],
        ["grado", "smallint", "NOT NULL, CHECK (grado IN (10, 11))", "Nivel académico."],
        ["anio_lectivo", "smallint", "NOT NULL", "Año académico (ej.: 2026)."],
        ["estado", "enum", "NOT NULL: activa, archivada", "Estado del curso."]
    ]
    add_table_data(doc, h_cols, d_clase, col_widths=[1.5, 1.2, 1.8, 2.5])

    # 4.12 evaluacion
    add_heading_with_spacing(doc, "4.12 evaluacion (RF-41 a RF-43)", level=3)
    add_p(doc, "Configuración de exámenes oficiales asignados a una clase, accesibles mediante código VF-XXXX.")
    d_eval = [
        ["id", "UUID", "PK", "Identificador."],
        ["clase_id", "UUID", "FK → clase, NOT NULL", "Clase destinataria."],
        ["titulo", "varchar(150)", "NOT NULL", "Nombre de la prueba."],
        ["codigo", "varchar(10)", "NOT NULL, UNIQUE", "Código alfanumérico único (ej. VF-4821)."],
        ["modo_seleccion", "enum", "NOT NULL: automatica, manual", "Generada por filtros o seleccionada pregunta por pregunta."],
        ["dificultad", "enum", "NULL: ponderada, facil, media, dificil", "Filtro en caso de selección automática."],
        ["cantidad_preguntas", "smallint", "NOT NULL, CHECK (cantidad BETWEEN 5 AND 50)", "Número total de preguntas."],
        ["duracion_min", "smallint", "NULL", "Límite en minutos. NULL indica sin límite de tiempo."],
        ["intentos_max", "smallint", "NOT NULL, default 1", "Máximo de intentos permitidos."],
        ["abre_en", "timestamptz", "NOT NULL", "Apertura de la ventana de presentación."],
        ["cierra_en", "timestamptz", "NOT NULL, CHECK (cierra_en > abre_en)", "Cierre de la ventana de presentación."],
        ["mostrar_resultado", "enum", "NOT NULL: inmediato, diferido", "Inmediato al entregar o tras cierre de la evaluación."],
        ["barajar", "boolean", "NOT NULL, default true", "Orden aleatorio de preguntas y opciones."],
        ["estado", "enum", "NOT NULL: borrador, publicada, cerrada, archivada", "Ciclo de vida del examen."],
        ["creada_por", "UUID", "FK → usuario, NOT NULL", "Docente autor."]
    ]
    add_table_data(doc, h_cols, d_eval, col_widths=[1.5, 1.2, 1.8, 2.5])

    # 4.13 evaluacion_tema y evaluacion_pregunta
    add_heading_with_spacing(doc, "4.13 evaluacion_tema y evaluacion_pregunta", level=3)
    add_p(doc, "Tablas de composición de la evaluación. 'evaluacion_pregunta' congela el conjunto de preguntas definitivas al momento de publicar.")
    d_eval_comp = [
        ["evaluacion_tema", "evaluacion_id (FK), tema_id (FK)", "PK compuesta", "Temas seleccionados para generación automática."],
        ["evaluacion_pregunta", "evaluacion_id (FK), pregunta_id (FK), orden (smallint), puntos (numeric(3,1) default 1.0)", "PK compuesta", "Preguntas fijadas para el examen oficial."]
    ]
    add_table_data(doc, ["Tabla", "Columnas", "Restricciones", "Descripción"], d_eval_comp, col_widths=[1.5, 2.2, 1.3, 2.0])

    # 4.14 intento
    add_heading_with_spacing(doc, "4.14 intento (RF-33, RF-35)", level=3)
    add_p(doc, "Registro de presentación de pruebas oficiales de clase. Solo almacena evaluaciones formales (no prácticas libres).")
    d_intento = [
        ["id", "UUID", "PK", "Identificador único."],
        ["estudiante_id", "UUID", "FK → usuario, NOT NULL", "Estudiante evaluado."],
        ["evaluacion_id", "UUID", "FK → evaluacion, NOT NULL", "Evaluación formal presentada."],
        ["numero_intento", "smallint", "NOT NULL, default 1", "Número ordinal del intento."],
        ["total_preguntas", "smallint", "NOT NULL", "Cantidad de preguntas asignadas."],
        ["estado", "enum", "NOT NULL: en_curso, finalizado, expirado, anulado", "Estado del intento."],
        ["inicio_en", "timestamptz", "NOT NULL, default now()", "Inicio cronometrado."],
        ["limite_en", "timestamptz", "NULL", "Hora máxima de entrega fijada por el servidor."],
        ["fin_en", "timestamptz", "NULL", "Hora de entrega efectiva."],
        ["aciertos", "smallint", "NULL", "Total de respuestas correctas logradas."],
        ["puntaje", "numeric(3,2)", "NULL, CHECK (puntaje BETWEEN 1.00 AND 5.00)", "Nota definitiva en escala colombiana de 1.00 a 5.00."]
    ]
    add_table_data(doc, h_cols, d_intento, col_widths=[1.5, 1.2, 1.8, 2.5])
    add_p(doc, "Restricción única: UNIQUE (estudiante_id, evaluacion_id, numero_intento). El servidor bloquea intentos adicionales que excedan 'intentos_max'.")

    # 4.15 intento_respuesta
    add_heading_with_spacing(doc, "4.15 intento_respuesta (RF-34)", level=3)
    add_p(doc, "Almacena las respuestas emitidas en tiempo real para salvaguardar el avance ante desconexiones.")
    d_resp = [
        ["id", "UUID", "PK", "Identificador."],
        ["intento_id", "UUID", "FK → intento, NOT NULL", "Intento formal en curso."],
        ["pregunta_id", "UUID", "FK → pregunta, NOT NULL", "Pregunta respondida."],
        ["orden", "smallint", "NOT NULL", "Posición mostrada al alumno."],
        ["opcion_elegida_id", "UUID", "FK → opcion_pregunta, NULL", "Selección en opción múltiple."],
        ["texto_respuesta", "varchar(200)", "NULL", "Texto ingresado en preguntas de completar."],
        ["es_correcta", "boolean", "NULL", "Calificación calculada en servidor al entregar."],
        ["respondida_en", "timestamptz", "NOT NULL, default now()", "Marca de guardado progresivo."]
    ]
    add_table_data(doc, h_cols, d_resp, col_widths=[1.5, 1.2, 1.8, 2.5])
    add_p(doc, "Restricción única: UNIQUE (intento_id, pregunta_id).")

    # 4.16 evento_examen
    add_heading_with_spacing(doc, "4.16 evento_examen (RF-36)", level=3)
    add_p(doc, "Telemetría de incidencias durante el Modo Examen. Indicadores informativos para el docente; no invalidan automáticamente la prueba.")
    d_evento = [
        ["id", "bigint", "PK autoincremental", "Identificador."],
        ["intento_id", "UUID", "FK → intento, NOT NULL", "Intento en el que ocurrió la alerta."],
        ["tipo", "enum", "NOT NULL: perdida_foco, salida_pestana, pantalla_completa_salida, reconexion", "Naturaleza de la señal."],
        ["ocurrido_en", "timestamptz", "NOT NULL, default now()", "Marca de tiempo del cliente/servidor."],
        ["detalle", "jsonb", "NULL", "Información adicional de telemetría (segundos fuera, agente de usuario)."]
    ]
    add_table_data(doc, h_cols, d_evento, col_widths=[1.5, 1.2, 1.8, 2.5])

    # 5. Relaciones principales
    add_heading_with_spacing(doc, "5. Relaciones principales", level=2)
    h_rel = ["Entidad origen", "Cardinalidad", "Entidad destino", "Regla de negocio / Integridad"]
    d_rel = [
        ["usuario (docente)", "1 a N", "clase", "Un docente gestiona múltiples clases."],
        ["clase", "1 a N", "usuario (estudiante)", "Un estudiante solo puede estar matriculado en una clase a la vez (clase_id directo)."],
        ["usuario", "1 a N", "consentimiento", "Historial de aceptaciones legales de política de datos (Ley 1581)."],
        ["usuario (docente)", "1 a 1", "solicitud_docente", "Registro de validación y aprobación por el Owner."],
        ["clase", "1 a N", "evaluacion", "Cada evaluación pertenece a una clase concreta."],
        ["evaluacion", "N a M", "tema", "A través de 'evaluacion_tema'."],
        ["evaluacion", "N a M", "pregunta", "A través de 'evaluacion_pregunta' (congelada al publicar)."],
        ["tema", "1 a N", "pregunta", "Banco organizado por contenido pedagógico."],
        ["pregunta", "1 a N", "opcion_pregunta", "Opciones para modalidad de selección múltiple."],
        ["pregunta", "1 a N", "respuesta_aceptada", "Textos válidos para modalidad de respuesta escrita/completar."],
        ["evaluacion", "1 a N", "intento", "Intentos formales de evaluación."],
        ["usuario (estudiante)", "1 a N", "intento", "Historial de pruebas presentadas por el estudiante."],
        ["intento", "1 a N", "intento_respuesta", "Respuestas individuales guardadas progresivamente."],
        ["intento", "1 a N", "evento_examen", "Eventos de pérdida de foco o navegación sospechosa."]
    ]
    add_table_data(doc, h_rel, d_rel, col_widths=[1.8, 1.1, 1.8, 2.3])

    # 6. Reglas de cálculo de calificaciones
    add_heading_with_spacing(doc, "6. Reglas de cálculo de calificaciones", level=2)
    add_p(doc, "La plataforma utiliza la escala estándar del Ministerio de Educación Nacional de Colombia (1.0 a 5.0) para las evaluaciones oficiales:")
    add_p(doc, "Puntaje = 1.0 + (aciertos / total_preguntas) * 4.0", bold_prefix="Fórmula oficial: ")
    add_bullet(doc, "Comportamiento en extremos", "0 aciertos = 1.00; 100% de aciertos = 5.00.")
    add_bullet(doc, "Aprobación", "La nota mínima aprobatoria estándar en el colegio es 3.00 (equivale al 50% de aciertos).")
    add_bullet(doc, "Redondeo", "El puntaje se guarda internamente con dos decimales numeric(3,2) y se redondea a un decimal en reportes académicos y planillas.")
    add_bullet(doc, "Cero gamificación", "La nota oficial no incluye bonificaciones por XP, rachas ni velocidad; refleja exclusivamente el desempeño académico.")

    # 7. Índices recomendados
    add_heading_with_spacing(doc, "7. Índices recomendados", level=2)
    h_idx = ["Tabla", "Definición del índice", "Objetivo de optimización"]
    d_idx = [
        ["usuario", "UNIQUE (lower(nickname))", "Autenticación rápida y garantía de unicidad sin distinguir mayúsculas."],
        ["usuario", "INDEX (clase_id) WHERE clase_id IS NOT NULL", "Listado inmediato de estudiantes por curso en la Suite del Docente."],
        ["evaluacion", "UNIQUE (codigo)", "Búsqueda instantánea de prueba al ingresar código VF-XXXX."],
        ["evaluacion", "INDEX (clase_id, estado)", "Filtrado de exámenes activos por clase."],
        ["intento", "INDEX (evaluacion_id, estado)", "Tablero de control en tiempo real del docente mientras rinden el examen."],
        ["intento", "INDEX (estudiante_id, inicio_en DESC)", "Consulta de historial de notas del estudiante."],
        ["intento_respuesta", "UNIQUE (intento_id, pregunta_id)", "Integridad y guardado progresivo sin duplicados."],
        ["pregunta", "INDEX (tema_id, dificultad, estado, tipo)", "Generación ágil de paquetes de examen y práctica."]
    ]
    add_table_data(doc, h_idx, d_idx, col_widths=[1.8, 2.7, 2.5])

    # 8. Seguridad y privacidad (Ley 1581 y Ley 1098)
    add_heading_with_spacing(doc, "8. Seguridad y privacidad en la base de datos", level=2)
    add_bullet(doc, "Protección de claves de respuesta (RNF-04)", "El campo 'opcion_pregunta.es_correcta' y la tabla 'respuesta_aceptada' nunca son serializados por la API durante exámenes oficiales. La corrección se efectúa estrictamente en el backend al entregar.")
    add_bullet(doc, "Minimización radical de datos de menores", "La base de datos no solicita número de documento de identidad, dirección residencial, teléfono ni fecha de nacimiento. Solo maneja nombre completo, nickname institucional y grado escolar.")
    add_bullet(doc, "Práctica libre privada", "Las sesiones de práctica libre no generan registros en la base de datos, evitando trazabilidad conductual innecesaria.")
    add_bullet(doc, "Aislamiento por docente", "Las consultas de resultados en la API validan que 'clase.docente_id == request.user.id', imposibilitando que un docente acceda a notas de cursos ajenos.")
    add_bullet(doc, "Baja de cuenta y anonimización (RF-07)", "Al solicitar borrado de cuenta, se disocian datos personales (nombre y nickname reemplazados por identificador anónimo) conservando únicamente la nota académica requerida por la institución educativa.")

    # 9. Trazabilidad con el SRS v0.2
    add_heading_with_spacing(doc, "9. Trazabilidad con el SRS v0.2", level=2)
    h_traz = ["Requisitos del SRS", "Tablas del modelo involucradas", "Propósito verificado"]
    d_traz = [
        ["RF-01 a RF-07", "usuario, consentimiento", "Registro mínimo, consentimiento Ley 1581 y gestión de cuentas."],
        ["RF-04, RF-50 a RF-53", "solicitud_docente, auditoria", "Flujo de aprobación de docentes y fiscalización del Owner."],
        ["RF-10 a RF-13", "tema, formula, audio_ejemplo", "Módulo de estudio y chuleta de fórmulas."],
        ["RF-20 a RF-24", "pregunta, opcion_pregunta, respuesta_aceptada", "Banco de preguntas para práctica y examen (opción múltiple y completar)."],
        ["RF-30 a RF-37", "evaluacion, intento, intento_respuesta, evento_examen", "Ejecución de exámenes oficiales y modo examen seguro."],
        ["RF-40 a RF-47", "clase, evaluacion, evaluacion_tema, evaluacion_pregunta, pregunta", "Suite del docente para cursos, exámenes y banco propio."]
    ]
    add_table_data(doc, h_traz, d_traz, col_widths=[1.8, 2.7, 2.5])

    # 10. Estado del diseño
    add_heading_with_spacing(doc, "10. Estado del diseño", level=2)
    add_p(doc, "El modelo de datos se encuentra CERRADO y aprobado en su versión 0.2. Todos los supuestos fueron resueltos en coordinación con los requerimientos pedagógicos de la Institución Educativa Promoción Social. Las siguientes fases son la implementación de migraciones en Django ORM y los controladores de la API REST.")
    
    doc.save(filepath)
    print(f"Modelo de datos guardado exitosamente en {filepath}")

# Run generator
os.makedirs("Docs", exist_ok=True)
generate_modelo_datos("Docs/VerbFlow — Modelo de Datos.docx")
